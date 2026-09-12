import csv
import io

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from reportlab.lib import colors
from reportlab.lib.pagesizes import landscape, A4
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet

COLUMNS = [
    ('post_date', 'Date'),
    ('narration', 'Description'),
    ('dr_cr', 'Type'),
    ('flag', 'Flag'),
    ('matched_ref', 'Ref No'),
    ('amount', 'Amount'),
    ('status', 'Status'),
    ('matched_date', 'Actual Date'),
    ('credit_amt', 'Credit Txn'),
    ('debit_amt', 'Debit Txn'),
    ('day', 'Day'),
]


def _cell(entry, key):
    val = entry.get(key)
    if val is None:
        return ''
    if key in ('amount', 'credit_amt', 'debit_amt'):
        try:
            return f'{float(val):,.2f}'
        except (TypeError, ValueError):
            return ''
    return str(val)


def export_csv(entries, gl_code, gl_desc, run_label):
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow([f'{gl_code} - {gl_desc}', run_label])
    writer.writerow([])
    writer.writerow([label for _, label in COLUMNS])
    for e in entries:
        writer.writerow([_cell(e, key) for key, _ in COLUMNS])
    return buf.getvalue().encode('utf-8-sig')  # BOM so Excel opens it with correct encoding


def export_xlsx(entries, gl_code, gl_desc, run_label):
    wb = Workbook()
    ws = wb.active
    ws.title = gl_code

    title_font = Font(name='Arial', bold=True, size=13, color='1F4E78')
    sub_font = Font(name='Arial', size=10, italic=True, color='555555')
    header_font = Font(name='Arial', bold=True, color='FFFFFF', size=10)
    header_fill = PatternFill('solid', fgColor='1F4E78')
    normal_font = Font(name='Arial', size=10)
    thin = Side(style='thin', color='BFBFBF')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    status_fills = {
        'Matched': PatternFill('solid', fgColor='C6EFCE'),
        'Pending - Multiple Possible Matches': PatternFill('solid', fgColor='FFC7CE'),
    }
    default_pending_fill = PatternFill('solid', fgColor='FFEB9C')
    na_fill = PatternFill('solid', fgColor='D9D9D9')
    excluded_fill = PatternFill('solid', fgColor='D9D2E9')

    ws['A1'] = f'{gl_code} - {gl_desc}'
    ws['A1'].font = title_font
    ws['A2'] = run_label
    ws['A2'].font = sub_font

    header_row = 4
    for c, (_, label) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=header_row, column=c, value=label)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = border
        cell.alignment = Alignment(horizontal='center', wrap_text=True)

    r = header_row + 1
    for e in entries:
        status = e.get('status', '')
        if status.startswith('Matched'):
            fill = status_fills['Matched']
        elif status.startswith('Excluded'):
            fill = excluded_fill
        elif 'Not Applicable' in status:
            fill = na_fill
        elif 'Multiple' in status:
            fill = status_fills['Pending - Multiple Possible Matches']
        else:
            fill = default_pending_fill

        for c, (key, _) in enumerate(COLUMNS, start=1):
            val = e.get(key)
            cell = ws.cell(row=r, column=c, value=val)
            cell.font = normal_font
            cell.border = border
            cell.fill = fill
            if key in ('amount', 'credit_amt', 'debit_amt') and val is not None:
                cell.number_format = '#,##0.00'
        r += 1

    widths = [11, 45, 6, 8, 20, 14, 30, 12, 14, 14, 6]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = f'A{header_row + 1}'

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def export_pdf(entries, gl_code, gl_desc, run_label):
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4),
                             leftMargin=12 * mm, rightMargin=12 * mm,
                             topMargin=12 * mm, bottomMargin=12 * mm)
    styles = getSampleStyleSheet()
    title_style = styles['Heading2']
    sub_style = styles['Normal']

    elements = [
        Paragraph(f'{gl_code} - {gl_desc}', title_style),
        Paragraph(run_label, sub_style),
        Spacer(1, 8),
    ]

    pdf_columns = [c for c in COLUMNS if c[0] != 'narration'] + [('narration', 'Description')]
    # keep narration readable by putting it last and letting the table wrap it,
    # rather than truncating other numeric columns to make room
    header = [label for _, label in pdf_columns]
    data = [header]
    para_style = styles['BodyText']
    para_style.fontSize = 7
    para_style.leading = 8

    for e in entries:
        row = []
        for key, _ in pdf_columns:
            text = _cell(e, key)
            if key in ('narration', 'status'):
                row.append(Paragraph(text, para_style))
            else:
                row.append(text)
        data.append(row)

    col_widths = [16 * mm, 10 * mm, 10 * mm, 18 * mm, 22 * mm, 42 * mm, 16 * mm, 22 * mm, 22 * mm, 8 * mm, 65 * mm]
    numeric_cols = [4, 7, 8]  # Amount, Credit Txn, Debit Txn
    table = Table(data, colWidths=col_widths, repeatRows=1)
    style_commands = [
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1F4E78')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTSIZE', (0, 0), (-1, 0), 7),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 1), (-1, -1), 7),
        ('GRID', (0, 0), (-1, -1), 0.4, colors.HexColor('#BFBFBF')),
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F5F5F5')]),
    ]
    for col in numeric_cols:
        style_commands.append(('ALIGN', (col, 1), (col, -1), 'RIGHT'))
    table.setStyle(TableStyle(style_commands))
    elements.append(table)
    doc.build(elements)
    return buf.getvalue()
